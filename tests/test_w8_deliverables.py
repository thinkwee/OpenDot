"""Deliverables: attach/make_* tool executors, the /api/deliverables router,
preview rendering, share links, and receipt creation from bus "step" events.
"""

from __future__ import annotations

import asyncio
import json

import pytest

from opendot import ext as ext_mod
from opendot.computer import computer_for
from opendot.db import db, new_id
from opendot.gatekeeper import DEFAULT_POLICY
from opendot.tools import EXEC, TOOLS, Ctx

ext_mod.load_all()
from opendot.ext import deliverables as D  # noqa: E402


@pytest.fixture(autouse=True)
def _clean_deliverables():
    with db.lock:
        db.conn.execute("DELETE FROM deliverables")
        db.conn.commit()
    yield


def _agent(name="Pip"):
    return db.insert("agents", id=new_id("ag_"), name=name, emoji="✨", color="#FFB38A",
                     role="r")


def _ctx(agent, thread="th_test", run="run_test"):
    return Ctx(agent=agent, thread_id=thread, run_id=run)


def _run(coro):
    return asyncio.run(coro)


# ---------------- tool registration ----------------
def test_tools_registered():
    for name in ("attach", "make_document", "make_spreadsheet", "make_slides", "publish_page"):
        assert name in TOOLS
        assert name in EXEC
        assert DEFAULT_POLICY.get(name) == "allow"


# ---------------- make_document ----------------
def test_make_document_docx_attaches_and_records_row():
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Report", markdown="# Hi\n\nBody **bold**.",
                                   format="docx"))
    assert "error" not in r
    assert r["kind"] == "document"
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row and row["agent_id"] == ag["id"] and row["thread_id"] == "th_test"
    assert row["path"].endswith(".docx")
    assert ctx.extra["attachments"][0]["id"] == r["id"]


def test_docx_turns_markdown_tables_into_real_tables(tmp_path):
    from docx import Document
    out = tmp_path / "t.docx"
    D._md_to_docx("## Overview\n\n| State | Count |\n|---|---:|\n| **Done** | 13 |\n"
                  "| a \\| b | 6 |\nAfter the table.", out)
    d = Document(out)
    assert len(d.tables) == 1
    rows = [[c.text for c in r.cells] for r in d.tables[0].rows]
    assert rows == [["State", "Count"], ["Done", "13"], ["a | b", "6"]]
    assert d.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold
    assert not any("|" in p.text for p in d.paragraphs)
    assert "After the table." in [p.text for p in d.paragraphs]


def test_make_document_md_format():
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Notes", markdown="hello world", format="md"))
    assert r["kind"] == "document"
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row["path"].endswith(".md")


def test_make_document_pdf_format():
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="PDF Report", markdown="# Title\n\nSome text.",
                                   format="pdf"))
    assert "error" not in r, r
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row["path"].endswith(".pdf")
    assert row["size"] > 0


# ---------------- make_spreadsheet ----------------
def test_make_spreadsheet_xlsx():
    ag = _agent()
    ctx = _ctx(ag)
    sheets = [{"name": "Data", "columns": ["A", "B"], "rows": [[1, 2], [3, 4]]}]
    r = _run(EXEC["make_spreadsheet"](ctx, title="Sheet", sheets=sheets, format="xlsx"))
    assert r["kind"] == "spreadsheet"
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row["path"].endswith(".xlsx")


def test_make_spreadsheet_csv():
    ag = _agent()
    ctx = _ctx(ag)
    sheets = [{"name": "Data", "columns": ["A", "B"], "rows": [[1, 2]]}]
    r = _run(EXEC["make_spreadsheet"](ctx, title="Sheet CSV", sheets=sheets, format="csv"))
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row["path"].endswith(".csv")
    content = open(row["path"]).read()
    assert "A,B" in content and "1,2" in content


# ---------------- make_slides ----------------
def test_make_slides_pptx():
    ag = _agent()
    ctx = _ctx(ag)
    slides = [{"title": "Slide 1", "bullets": ["a", "b"], "notes": "speak"}]
    r = _run(EXEC["make_slides"](ctx, title="Deck", slides=slides))
    assert r["kind"] == "slides"
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row["path"].endswith(".pptx")


# ---------------- attach ----------------
def test_attach_copies_file_from_home():
    ag = _agent()
    c = computer_for(ag["id"])
    c.write_file("note.txt", "hello")
    ctx = _ctx(ag)
    r = _run(EXEC["attach"](ctx, path="note.txt"))
    assert "error" not in r
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    assert row["path"] != str(c.home / "note.txt")  # snapshotted, not the original
    assert open(row["path"]).read() == "hello"


def test_attach_rejects_path_escape():
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["attach"](ctx, path="/etc/passwd"))
    assert "error" in r
    assert "outside" in r["error"]


def test_attach_missing_file_errors():
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["attach"](ctx, path="does-not-exist.txt"))
    assert "error" in r


def test_attach_snapshot_survives_source_edit():
    """Editing the source file after attach must not change the delivered copy."""
    ag = _agent()
    c = computer_for(ag["id"])
    c.write_file("report.md", "version 1")
    ctx = _ctx(ag)
    r = _run(EXEC["attach"](ctx, path="report.md"))
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    c.write_file("report.md", "version 2 — edited later")
    assert open(row["path"]).read() == "version 1"


# ---------------- publish_page (kept working + now a deliverable) ----------------
def test_publish_page_still_writes_html_and_returns_url():
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["publish_page"](ctx, slug="my-app", html="<h1>Hi</h1>", title="My App"))
    assert r["url"] == f"/pages/{ag['id']}/my-app"
    from opendot.config import settings
    p = settings.DATA_DIR / "pages" / ag["id"] / "my-app.html"
    assert p.read_text() == "<h1>Hi</h1>"


def test_publish_page_also_creates_app_deliverable():
    ag = _agent()
    ctx = _ctx(ag)
    _run(EXEC["publish_page"](ctx, slug="dash", html="<p>x</p>", title="Dashboard"))
    row = db.one("SELECT * FROM deliverables WHERE agent_id=? AND kind='app'", ag["id"])
    assert row is not None
    assert row["title"] == "Dashboard"
    assert ctx.extra["attachments"][0]["kind"] == "app"


# ---------------- router ----------------
def test_router_list_and_get(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag, thread="th_router")
    _run(EXEC["make_document"](ctx, title="R1", markdown="hi", format="md"))
    resp = client.get("/api/deliverables", headers=auth_headers)
    assert resp.status_code == 200
    rows = resp.json()
    assert len(rows) == 1
    assert rows[0]["url"].endswith("/file")
    assert rows[0]["preview_url"].endswith("/preview")

    did = rows[0]["id"]
    resp2 = client.get(f"/api/deliverables/{did}", headers=auth_headers)
    assert resp2.status_code == 200
    assert resp2.json()["id"] == did


def test_router_filters_by_kind_and_thread(client, auth_headers):
    ag = _agent()
    ctx1 = _ctx(ag, thread="th_a")
    ctx2 = _ctx(ag, thread="th_b")
    _run(EXEC["make_document"](ctx1, title="Doc", markdown="x", format="md"))
    _run(EXEC["make_spreadsheet"](ctx2, title="Sheet",
                                  sheets=[{"columns": ["a"], "rows": [[1]]}], format="csv"))
    only_docs = client.get("/api/deliverables?kind=document", headers=auth_headers).json()
    assert all(r["kind"] == "document" for r in only_docs)
    only_b = client.get("/api/deliverables?thread=th_b", headers=auth_headers).json()
    assert len(only_b) == 1 and only_b[0]["kind"] == "spreadsheet"


def test_router_requires_auth(client):
    resp = client.get("/api/deliverables")
    assert resp.status_code == 401


def test_download_endpoint_sets_content_disposition(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Downloadable", markdown="hi", format="md"))
    resp = client.get(f"/api/deliverables/{r['id']}/file", headers=auth_headers)
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]


@pytest.mark.parametrize("maker,kw,expect_in_body", [
    ("make_document", dict(title="MD Prev", markdown="# Heading\nbody", format="md"), "Heading"),
    ("make_spreadsheet", dict(title="XL Prev",
                              sheets=[{"columns": ["A"], "rows": [["v1"]]}], format="xlsx"),
     "v1"),
    ("make_spreadsheet", dict(title="CSV Prev",
                              sheets=[{"columns": ["A"], "rows": [["v2"]]}], format="csv"),
     "v2"),
    ("make_slides", dict(title="Deck Prev", slides=[{"title": "S1", "bullets": ["b1"]}]), "b1"),
])
def test_preview_endpoint_renders_html(client, auth_headers, maker, kw, expect_in_body):
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC[maker](ctx, **kw))
    resp = client.get(f"/api/deliverables/{r['id']}/preview", headers=auth_headers)
    assert resp.status_code == 200
    assert "text/html" in resp.headers["content-type"]
    assert expect_in_body in resp.text


def test_preview_docx_via_mammoth(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Docx Prev", markdown="# Big Heading\nbody text",
                                   format="docx"))
    resp = client.get(f"/api/deliverables/{r['id']}/preview", headers=auth_headers)
    assert resp.status_code == 200
    assert "Big Heading" in resp.text


def test_preview_image_is_inline(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag)
    computer_for(ag["id"]).resolve("dot.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    r = _run(EXEC["attach"](ctx, path="dot.png", title="Img Prev"))
    resp = client.get(f"/api/deliverables/{r['id']}/preview", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("image/")


def test_preview_pdf_is_inline_pdf(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Pdf Prev", markdown="hello", format="pdf"))
    resp = client.get(f"/api/deliverables/{r['id']}/preview", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.headers["content-disposition"] == "inline"


def test_preview_missing_file_404s(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Vanishing", markdown="hi", format="md"))
    row = db.one("SELECT * FROM deliverables WHERE id=?", r["id"])
    import os
    os.remove(row["path"])
    resp = client.get(f"/api/deliverables/{r['id']}/preview", headers=auth_headers)
    assert resp.status_code == 404


# ---------------- share links ----------------
def test_share_then_public_fetch_then_revoke(client, auth_headers):
    ag = _agent()
    ctx = _ctx(ag)
    r = _run(EXEC["make_document"](ctx, title="Shareable", markdown="secret-ish content",
                                   format="md"))
    resp = client.post(f"/api/deliverables/{r['id']}/share", headers=auth_headers)
    assert resp.status_code == 200
    share_url = resp.json()["url"]
    assert share_url.startswith("/hook/")

    # public fetch needs no auth headers at all
    pub = client.get(share_url)
    assert pub.status_code == 200
    assert "secret-ish content" in pub.text

    # wrong hook token is rejected
    bad_url = share_url.replace(share_url.split("/")[2], "wrong-token")
    assert client.get(bad_url).status_code == 401

    # revoke
    rv = client.delete(f"/api/deliverables/{r['id']}/share", headers=auth_headers)
    assert rv.status_code == 200
    assert client.get(share_url).status_code == 404


def test_share_unknown_id_404s(client, auth_headers):
    resp = client.post("/api/deliverables/dl_does_not_exist/share", headers=auth_headers)
    assert resp.status_code == 404


# ---------------- receipts ----------------
@pytest.mark.parametrize("tool,args,expect_in_title", [
    ("send_email", {"to": "friend@example.com", "subject": "hi"}, "friend@example.com"),
    ("send_sms", {"to": "+15551234"}, "+15551234"),
    ("create_event", {"title": "Standup"}, "Standup"),
    ("mcp__icloud-calendar__create_event", {"title": "Dentist"}, "Dentist"),
    ("device_notify", {}, "device_notify".replace("device_", "")),
])
def test_receipt_created_from_step_event(tool, args, expect_in_title):
    ag = _agent()
    step = {"id": new_id("st_"), "thread_id": "th_rcpt", "run_id": "run_rcpt",
           "agent_id": ag["id"], "tool": tool, "args": args, "status": "done",
           "result": json.dumps({"ok": True})}
    D._on_bus("step", {"step": step})
    row = db.one("SELECT * FROM deliverables WHERE kind='receipt' AND source=?", tool)
    assert row is not None
    assert expect_in_title in row["title"]
    assert row["meta"]["args"] == args


def test_receipt_ignores_non_done_or_unmatched_tools():
    ag = _agent()
    before = db.one("SELECT count(*) n FROM deliverables WHERE kind='receipt'")["n"]
    D._on_bus("step", {"step": {"tool": "send_email", "status": "running",
                                "agent_id": ag["id"]}})
    D._on_bus("step", {"step": {"tool": "web_search", "status": "done", "agent_id": ag["id"]}})
    after = db.one("SELECT count(*) n FROM deliverables WHERE kind='receipt'")["n"]
    assert after == before


def test_receipt_ignores_non_step_events():
    before = db.one("SELECT count(*) n FROM deliverables")["n"]
    D._on_bus("message", {"message": {}})
    after = db.one("SELECT count(*) n FROM deliverables")["n"]
    assert after == before
