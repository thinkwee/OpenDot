"""Deliverables — real files as attachments, not just HTML Pages.

Like a good assistant: an agent finishes a task and hands back a *thing* — a report,
a spreadsheet, a slide deck, a chart — that shows up as a card in the chat
with Preview / Download / Share, plus a per-task Library. `publish_page`
still exists for interactive apps/dashboards, but it now also registers
itself here so it shows up next to everything else.

Contract with W1 (see docs/PLAN_V02.md "extension contract" + the W8 brief):
  - a tool executor stores a `deliverables` row and appends
    ``{"id","title","kind","mime","size","url","preview_url","agent_id"}``
    to ``ctx.extra.setdefault("attachments", [])``
  - W1's runtime copies ``ctx.extra["attachments"]`` into the final message's
    ``meta.attachments``; W1's ChatView renders each with
    ``import DeliverableCard from '../components/Deliverable'``
  - ``bus.emit("deliverable", deliverable=row)`` fires for every new file so
    the Library view can live-update.

Kinds: document, spreadsheet, slides, image, app, file, receipt.
"""

from __future__ import annotations

import csv
import hashlib
import hmac
import json
import logging
import mimetypes
import re
import shutil
from html import escape
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse

from ..bus import bus
from ..computer import computer_for
from ..config import settings
from ..db import db, new_id
from ..tools import TOOLS as CORE_TOOLS
from ..tools import Ctx, fn, register_tool

log = logging.getLogger("opendot.deliverables")

db.ensure_schema(
    """
    CREATE TABLE IF NOT EXISTS deliverables (
      id TEXT PRIMARY KEY, agent_id TEXT, thread_id TEXT, run_id TEXT, title TEXT,
      kind TEXT, path TEXT, mime TEXT, size INTEGER, sha TEXT, created REAL,
      source TEXT, meta TEXT DEFAULT '{}', share_id TEXT, share_enabled INTEGER DEFAULT 0
    );
    CREATE INDEX IF NOT EXISTS ix_deliverables_thread ON deliverables(thread_id, created);
    CREATE INDEX IF NOT EXISTS ix_deliverables_agent ON deliverables(agent_id, created);
    CREATE INDEX IF NOT EXISTS ix_deliverables_share ON deliverables(share_id);
    """
)

PROMPT = ("# Deliverables\nEverything you hand over is a file card: reports→make_document, "
          "data→make_spreadsheet, decks→make_slides, other files→attach, "
          "interactive mini-apps/dashboards→publish_page (a web-page file).\n"
          "Charts and pictures go inside your message: make_chart draws one right in your "
          "reply (no file), and so does a ```chart block you write yourself, holding the same "
          "JSON. A picture: ![caption](path on your computer, or an https link). Tables are "
          "Markdown tables. Most replies need none of these; use one when it says more than the "
          "sentence it replaces, and keep it to one or two.")

PALETTE = ["FFB38A", "8FD6B8", "B9A6F2", "8EC5FF", "FFD37A", "FF9EC4"]
S = {"type": "string"}
_EXT_KIND = {
    ".docx": "document", ".pdf": "document", ".md": "document", ".txt": "document",
    ".xlsx": "spreadsheet", ".csv": "spreadsheet", ".xls": "spreadsheet",
    ".pptx": "slides",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image", ".webp": "image",
    ".svg": "image",
    ".html": "app", ".htm": "app",
}


def _slug(title: str) -> str:
    return re.sub(r"[^a-z0-9\-]", "-", (title or "").lower()).strip("-")[:48] or new_id()


def _dir(did: str) -> Path:
    d = settings.DATA_DIR / "deliverables" / did
    d.mkdir(parents=True, exist_ok=True)
    return d


def _sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def _kind_from_ext(ext: str) -> str:
    return _EXT_KIND.get(ext.lower(), "file")


def _finish(ctx: Ctx, did: str, title: str, kind: str, path: Path, source: str,
            meta: dict | None = None) -> dict:
    """Record a deliverable row + attach it to the current run's reply."""
    size = path.stat().st_size
    sha = _sha256(path)
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    row = db.insert("deliverables", id=did, agent_id=ctx.agent["id"], thread_id=ctx.thread_id,
                    run_id=ctx.run_id, title=title, kind=kind, path=str(path), mime=mime,
                    size=size, sha=sha, source=source, meta=meta or {}, share_id=None,
                    share_enabled=0)
    att = {"id": did, "title": title, "kind": kind, "mime": mime, "size": size,
           "url": f"/api/deliverables/{did}/file",
           "preview_url": f"/api/deliverables/{did}/preview", "agent_id": ctx.agent["id"]}
    ctx.extra.setdefault("attachments", []).append(att)
    bus.emit("deliverable", deliverable=row)
    return {**att, "note": f"Delivered as a {kind} and already attached to your reply — do NOT "
                           "call attach on it again. Don't paste its full contents, just give "
                           "a short summary."}


# ---------------- attach ----------------
async def _attach(ctx: Ctx, path: str, title: str = "", kind: str = "") -> dict:
    c = computer_for(ctx.agent["id"])
    try:
        src = c.resolve(path)
    except PermissionError as e:
        return {"error": str(e)}
    if not src.is_file():
        return {"error": f"no such file: {path}"}
    did = new_id("dl_")
    dest = _dir(did) / src.name
    shutil.copy2(src, dest)
    return _finish(ctx, did, title or dest.stem, kind or _kind_from_ext(dest.suffix), dest,
                  "attach")


# ---------------- documents ----------------
def _add_runs(p, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 1:
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            p.add_run(part)


def _md_to_docx(markdown_text: str, out: Path, title: str = "") -> None:
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    rows: list[str] = []  # a markdown table being read, added when it ends
    for line in markdown_text.splitlines() + [""]:
        stripped = line.strip()
        if stripped.startswith("|"):
            rows.append(stripped)
            continue
        if rows:
            _add_table(doc, rows)
            rows = []
        if not stripped:
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif re.match(r"^[-*]\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith(">"):
            doc.add_paragraph(stripped.lstrip("> ").strip())
        else:
            _add_runs(doc.add_paragraph(), stripped)
    doc.save(out)


def _cells(row: str) -> list[str]:
    return [c.strip() for c in re.split(r"(?<!\\)\|", row.strip().strip("|"))]


def _add_table(doc, rows: list[str]) -> None:
    body = [_cells(r) for r in rows if not re.fullmatch(r"[|:\-\s]+", r)]
    if not body:
        return
    width = max(len(r) for r in body)
    table = doc.add_table(rows=len(body), cols=width)
    table.style = "Table Grid"
    for i, cells in enumerate(body):
        for j in range(width):
            p = table.cell(i, j).paragraphs[0]
            text = cells[j].replace("\\|", "|") if j < len(cells) else ""
            if i == 0 and len(rows) > 1 and re.fullmatch(r"[|:\-\s]+", rows[1]):
                p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", text)).bold = True  # header row
            else:
                _add_runs(p, text)
    doc.add_paragraph()


def _wrap_html(body: str, title: str = "") -> str:
    return f"""<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title or 'Preview')}</title>
<style>
  body {{ font-family: Nunito, ui-rounded, system-ui, -apple-system, sans-serif; color:#2b2233;
         background:#fffdfb; margin:0; padding:24px; line-height:1.6; }}
  h1,h2,h3 {{ color:#ff8a5c; }}
  table {{ border-collapse:collapse; width:100%; margin:12px 0; }}
  th, td {{ border:1px solid #f1e3d8; padding:6px 10px; text-align:left; font-size:14px; }}
  th {{ background:#fff3ea; }}
  code, pre {{ background:#fff3ea; border-radius:8px; padding:2px 6px; }}
  pre {{ padding:12px; overflow:auto; white-space:pre-wrap; }}
  .slide-card {{ border:1px solid #f1e3d8; border-radius:14px; padding:16px 18px; margin:12px 0;
                background:#fff; box-shadow:0 2px 10px rgba(123,72,40,.08); }}
  .slide-card h3 {{ margin-top:0; }}
  ul {{ margin:6px 0; padding-left:22px; }}
  .notes {{ color:#6f6477; font-size:13px; margin-top:8px; }}
  .muted {{ color:#6f6477; }}
  img {{ max-width:100%; }}
</style></head><body>{body}</body></html>"""


async def _md_to_pdf(markdown_text: str, out: Path, title: str = "") -> None:
    import markdown as mdlib
    from playwright.async_api import async_playwright

    body = mdlib.markdown(markdown_text, extensions=["tables", "fenced_code"])
    if title:
        body = f"<h1>{escape(title)}</h1>" + body
    html = _wrap_html(body, title)
    from ..computer import _find_chromium
    exe = settings.CHROMIUM_PATH or _find_chromium()
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(executable_path=exe or None)
        try:
            page = await browser.new_page()
            await page.set_content(html, wait_until="load")
            await page.pdf(path=str(out), format="A4", print_background=True,
                           margin={"top": "18mm", "bottom": "18mm", "left": "16mm",
                                  "right": "16mm"})
        finally:
            await browser.close()


async def _make_document(ctx: Ctx, title: str, markdown: str, format: str = "docx") -> dict:
    did = new_id("dl_")
    d = _dir(did)
    fmt = (format or "docx").lower()
    if fmt == "md":
        p = d / f"{_slug(title)}.md"
        p.write_text(markdown)
    elif fmt == "pdf":
        p = d / f"{_slug(title)}.pdf"
        try:
            await _md_to_pdf(markdown, p, title)
        except Exception as e:
            log.exception("pdf render failed")
            return {"error": f"could not render pdf: {e}"}
    else:
        p = d / f"{_slug(title)}.docx"
        _md_to_docx(markdown, p, title)
    return _finish(ctx, did, title, "document", p, "make_document")


# ---------------- spreadsheets ----------------
def _write_xlsx(p: Path, sheets: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="FFE3D0")
    header_font = Font(bold=True, color="2B2233")
    for sh in sheets or [{"name": "Sheet1", "columns": [], "rows": []}]:
        ws = wb.create_sheet(title=(sh.get("name") or "Sheet1")[:31] or "Sheet1")
        cols = sh.get("columns") or []
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left")
        for ri, row in enumerate(sh.get("rows") or [], 2):
            for ci, v in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=v)
        for ci, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = max(
                10, min(40, len(str(c)) + 4))
    wb.save(p)


async def _make_spreadsheet(ctx: Ctx, title: str, sheets: list[dict],
                            format: str = "xlsx") -> dict:
    did = new_id("dl_")
    d = _dir(did)
    fmt = (format or "xlsx").lower()
    if fmt == "csv":
        p = d / f"{_slug(title)}.csv"
        sh = (sheets or [{}])[0]
        with open(p, "w", newline="") as f:
            w = csv.writer(f)
            if sh.get("columns"):
                w.writerow(sh["columns"])
            for row in sh.get("rows") or []:
                w.writerow(row)
    else:
        p = d / f"{_slug(title)}.xlsx"
        _write_xlsx(p, sheets)
    return _finish(ctx, did, title, "spreadsheet", p, "make_spreadsheet")


# ---------------- slides ----------------
def _build_pptx(ctx: Ctx, deck_title: str, slides: list[dict], out: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    agent = ctx.agent
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def bg(slide, hexcolor):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(hexcolor)

    s = prs.slides.add_slide(blank)
    bg(s, "2B2233")
    tb = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.3), Inches(1.6))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = deck_title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string("FFFFFF")
    sub = s.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.3), Inches(0.6))
    sp = sub.text_frame.paragraphs[0]
    sp.text = f"{agent.get('emoji', '✨')} {agent.get('name', 'Pip')}"
    sp.font.size = Pt(20)
    sp.font.color.rgb = RGBColor.from_string(PALETTE[0])

    for i, sl in enumerate(slides):
        color = PALETTE[i % len(PALETTE)]
        s = prs.slides.add_slide(blank)
        bg(s, "FFFDFB")
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333),
                                 Inches(1.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(color)
        bar.line.fill.background()
        bar.text_frame.word_wrap = True
        bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        tp = bar.text_frame.paragraphs[0]
        tp.text = sl.get("title") or f"Slide {i + 1}"
        tp.font.size = Pt(30)
        tp.font.bold = True
        tp.font.color.rgb = RGBColor.from_string("2B2233")

        has_img = bool(sl.get("image_path"))
        body = s.shapes.add_textbox(Inches(1), Inches(1.7),
                                    Inches(8.0 if has_img else 11.3), Inches(5.3))
        body.text_frame.word_wrap = True
        bullets = sl.get("bullets") or []
        if bullets:
            body.text_frame.paragraphs[0].text = bullets[0]
            body.text_frame.paragraphs[0].font.size = Pt(20)
            body.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("2B2233")
            for b in bullets[1:]:
                para = body.text_frame.add_paragraph()
                para.text = b
                para.font.size = Pt(20)
                para.font.color.rgb = RGBColor.from_string("2B2233")
        if has_img:
            try:
                real = computer_for(agent["id"]).resolve(sl["image_path"])
                s.shapes.add_picture(str(real), Inches(9.2), Inches(1.7), height=Inches(4.5))
            except Exception:
                log.debug("slide image skipped: %s", sl.get("image_path"))
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]

    prs.save(out)


async def _make_slides(ctx: Ctx, title: str, slides: list[dict]) -> dict:
    did = new_id("dl_")
    p = _dir(did) / f"{_slug(title)}.pptx"
    _build_pptx(ctx, title, slides or [], p)
    return _finish(ctx, did, title, "slides", p, "make_slides")


# ---------------- charts ----------------
async def _make_chart(ctx: Ctx, title: str = "", type: str = "bar", data: dict | str | None = None,
                      **_) -> dict:
    """The chart is drawn inside the reply (a ```chart block the chat renders), not a file."""
    if isinstance(data, str):
        try:
            data = json.loads(data)
        except ValueError:
            return {"error": "data must be {labels:[...], series:[{name, values:[...]}]}"}
    data = data or {}
    labels = [str(x) for x in data.get("labels") or []]
    series = data.get("series") or ([{"name": title, "values": data["values"]}]
                                    if data.get("values") else [])
    if not labels or not series:
        return {"error": "data needs labels and at least one series of values"}
    spec = {"type": type if type in ("bar", "line", "pie") else "bar", "title": title,
            "labels": labels, "series": [{"name": str(x.get("name", "")),
                                          "values": list(x.get("values") or [])} for x in series]}
    if isinstance(data.get("y"), dict):
        spec["y"] = {k: v for k, v in data["y"].items() if k in ("min", "max")}
    ctx.extra.setdefault("charts", []).append(
        "```chart\n" + json.dumps(spec, ensure_ascii=False) + "\n```")
    return {"ok": True, "note": "It will show inside your reply, under your text. Don't repeat "
                                "the numbers; say what they mean."}


# ---------------- publish_page (kept, now also a deliverable) ----------------
async def _publish_page(ctx: Ctx, slug: str, html: str, title: str = "") -> dict:
    slug2 = _slug(slug)
    pdir = settings.DATA_DIR / "pages" / ctx.agent["id"]
    pdir.mkdir(parents=True, exist_ok=True)
    fpath = pdir / f"{slug2}.html"
    fpath.write_text(html)
    url = f"/pages/{ctx.agent['id']}/{slug2}"
    bus.emit("page", agent_id=ctx.agent["id"], url=url, title=title or slug2)
    did = new_id("dl_")
    size = fpath.stat().st_size
    row = db.insert("deliverables", id=did, agent_id=ctx.agent["id"], thread_id=ctx.thread_id,
                    run_id=ctx.run_id, title=title or slug2, kind="app", path=str(fpath),
                    mime="text/html", size=size, sha=_sha256(fpath), source="publish_page",
                    meta={"url": url}, share_id=None, share_enabled=0)
    att = {"id": did, "title": title or slug2, "kind": "app", "mime": "text/html", "size": size,
          "url": url, "preview_url": url, "agent_id": ctx.agent["id"]}
    ctx.extra.setdefault("attachments", []).append(att)
    bus.emit("deliverable", deliverable=row)
    return {"url": url, "note": "Created as a web-page file and already attached to your reply as a card — don't attach it again."}


# ---------------- receipts ----------------
# create_event: an event a calendar app added (mcp__<app>__create_event)
_RECEIPT_RE = re.compile(r"^(send_email|send_sms|(mcp__.+__)?create_event|device_.+)$")


def _receipt_title(tool: str, args: dict) -> str:
    if tool == "send_email":
        return f"✉️ Email sent to {args.get('to', '?')}"
    if tool == "send_sms":
        return f"📱 SMS sent to {args.get('to', '?')}"
    if tool.endswith("create_event"):
        return f"📅 Event created: {args.get('title') or args.get('summary') or '(untitled)'}"
    if tool.startswith("device_"):
        return f"🖥️ Device action: {tool.replace('device_', '')}"
    return f"✅ {tool}"


def _on_bus(kind: str, data: dict) -> None:
    if kind != "step":
        return
    step = data.get("step") or {}
    if step.get("status") != "done" or not _RECEIPT_RE.match(step.get("tool", "")):
        return
    try:
        args = step.get("args") or {}
        tool = step["tool"]
        title = _receipt_title(tool, args)
        did = new_id("dl_")
        d = _dir(did)
        meta = {"tool": tool, "args": args}
        import json
        try:
            meta["result"] = json.loads(step.get("result") or "{}")
        except Exception:
            meta["result"] = {"raw": step.get("result")}
        p = d / "receipt.json"
        p.write_text(json.dumps(meta, indent=2, default=str))
        row = db.insert("deliverables", id=did, agent_id=step.get("agent_id"),
                        thread_id=step.get("thread_id"), run_id=step.get("run_id"), title=title,
                        kind="receipt", path=str(p), mime="application/json", size=p.stat().st_size,
                        sha=_sha256(p), source=tool, meta=meta, share_id=None, share_enabled=0)
        bus.emit("deliverable", deliverable=row)
    except Exception:
        log.exception("receipt creation failed")


bus.listen(_on_bus)


# ---------------- preview rendering ----------------
def _xlsx_table(p: Path, max_rows: int = 200) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    n_sheets = len(wb.sheetnames)
    names = ", ".join(wb.sheetnames)
    wb.close()
    if not rows:
        return "<p class='muted'>Empty sheet.</p>"
    head, *body = rows
    out = []
    if n_sheets > 1:
        out.append(f"<p class='muted'>Sheet 1 of {n_sheets}: {escape(names)}</p>")
    out.append("<table><thead><tr>")
    out += [f"<th>{escape('' if c is None else str(c))}</th>" for c in head]
    out.append("</tr></thead><tbody>")
    for r in body[:max_rows]:
        out.append("<tr>" + "".join(
            f"<td>{escape('' if c is None else str(c))}</td>" for c in r) + "</tr>")
    out.append("</tbody></table>")
    if len(body) > max_rows:
        out.append(f"<p class='muted'>… {len(body) - max_rows} more rows</p>")
    return "".join(out)


def _csv_table(p: Path, max_rows: int = 200) -> str:
    with open(p, newline="", errors="replace") as f:
        rows = list(csv.reader(f))
    if not rows:
        return "<p class='muted'>Empty file.</p>"
    head, *body = rows
    out = ["<table><thead><tr>"]
    out += [f"<th>{escape(c)}</th>" for c in head]
    out.append("</tr></thead><tbody>")
    for r in body[:max_rows]:
        out.append("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in r) + "</tr>")
    out.append("</tbody></table>")
    if len(body) > max_rows:
        out.append(f"<p class='muted'>… {len(body) - max_rows} more rows</p>")
    return "".join(out)


def _pptx_cards(p: Path) -> str:
    from pptx import Presentation

    prs = Presentation(p)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title, bullets = "", []
        real_title = slide.shapes.title
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            txt = shape.text_frame.text.strip()
            if not txt:
                continue
            # a real title placeholder wins; otherwise the first text shape on the
            # slide is treated as the title (our own make_slides template uses a
            # plain textbox for the header, not a placeholder)
            if (real_title is not None and shape == real_title and not title) or \
                    (real_title is None and not title and not bullets):
                title = txt
            else:
                bullets.extend(l for l in txt.splitlines() if l.strip())
        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        out.append(f"<div class='slide-card'><h3>{i}. {escape(title or '(untitled)')}</h3>"
                  "<ul>" + "".join(f"<li>{escape(b)}</li>" for b in bullets) + "</ul>" +
                  (f"<div class='notes'>📝 {escape(notes)}</div>" if notes else "") + "</div>")
    return "".join(out) or "<p class='muted'>No slides.</p>"


def _render_preview(row: dict):
    """Return a Starlette Response for previewing a deliverable's file."""
    p = Path(row["path"])
    if not p.exists():
        raise HTTPException(404, "file missing on disk")
    if row["kind"] == "receipt":
        return JSONResponse(row.get("meta") or {})
    ext = p.suffix.lower()
    if ext == ".pdf":
        return FileResponse(p, media_type="application/pdf",
                            headers={"Content-Disposition": "inline"})
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"):
        return FileResponse(p)
    if ext in (".html", ".htm"):
        return HTMLResponse(p.read_text(errors="replace"), headers={
            "Content-Security-Policy": "default-src 'self' 'unsafe-inline' data: blob: "
                                       "https://cdn.jsdelivr.net https://cdnjs.cloudflare.com; "
                                       "img-src * data: blob:; connect-src 'self'"})
    try:
        if ext == ".md":
            import markdown as mdlib
            return HTMLResponse(_wrap_html(
                mdlib.markdown(p.read_text(errors="replace"), extensions=["tables", "fenced_code"]),
                row["title"]))
        if ext == ".docx":
            import mammoth
            with open(p, "rb") as f:
                result = mammoth.convert_to_html(f)
            return HTMLResponse(_wrap_html(result.value, row["title"]))
        if ext == ".xlsx":
            return HTMLResponse(_wrap_html(_xlsx_table(p), row["title"]))
        if ext == ".csv":
            return HTMLResponse(_wrap_html(_csv_table(p), row["title"]))
        if ext == ".pptx":
            return HTMLResponse(_wrap_html(_pptx_cards(p), row["title"]))
    except Exception as e:
        log.exception("preview render failed for %s", p)
        return HTMLResponse(_wrap_html(f"<p class='muted'>Preview failed: {escape(str(e))}</p>"))
    try:
        text = p.read_text(errors="replace")[:20000]
        return HTMLResponse(_wrap_html(f"<pre>{escape(text)}</pre>", row["title"]))
    except Exception:
        raise HTTPException(415, "no preview available for this file type")


# ---------------- router ----------------
router = APIRouter()


def _hook_token() -> str:
    import secrets
    t = db.kv_get("hook_token")
    if not t:
        t = secrets.token_urlsafe(16)
        db.kv_set("hook_token", t)
    return t


def _augment(row: dict) -> dict:
    """The DB row only stores the on-disk path; add the URLs the frontend needs."""
    row = dict(row)
    row["url"] = f"/api/deliverables/{row['id']}/file"
    row["preview_url"] = f"/api/deliverables/{row['id']}/preview"
    return row


@router.get("/api/deliverables")
async def list_deliverables(thread: str = "", agent: str = "", kind: str = "",
                            limit: int = 500):
    q = "SELECT * FROM deliverables WHERE 1=1"
    args: list = []
    if thread:
        q += " AND thread_id=?"
        args.append(thread)
    if agent:
        q += " AND agent_id=?"
        args.append(agent)
    if kind:
        q += " AND kind=?"
        args.append(kind)
    q += " ORDER BY created DESC LIMIT ?"
    args.append(limit)
    return [_augment(r) for r in db.q(q, *args)]


@router.get("/api/deliverables/{did}")
async def get_deliverable(did: str):
    row = db.one("SELECT * FROM deliverables WHERE id=?", did)
    if not row:
        raise HTTPException(404)
    return _augment(row)


@router.get("/api/deliverables/{did}/file")
async def download(did: str):
    row = db.one("SELECT * FROM deliverables WHERE id=?", did)
    if not row:
        raise HTTPException(404)
    p = Path(row["path"])
    if not p.exists():
        raise HTTPException(404, "file missing on disk")
    # keep the real title (中文 included); only drop characters file systems reject
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "-", row["title"] or "").strip(" .-")[:80]
    return FileResponse(p, filename=f"{name or _slug(row['title'])}{p.suffix}",
                        media_type=row["mime"] or None)


@router.get("/api/deliverables/{did}/preview")
async def preview(did: str):
    row = db.one("SELECT * FROM deliverables WHERE id=?", did)
    if not row:
        raise HTTPException(404)
    return _render_preview(row)


@router.post("/api/deliverables/{did}/share")
async def share(did: str):
    row = db.one("SELECT * FROM deliverables WHERE id=?", did)
    if not row:
        raise HTTPException(404)
    sid = row.get("share_id") or new_id("sh_")
    db.update("deliverables", did, share_id=sid, share_enabled=1)
    return {"share_id": sid, "url": f"/hook/{_hook_token()}/d/{sid}"}


@router.delete("/api/deliverables/{did}/share")
async def unshare(did: str):
    row = db.one("SELECT * FROM deliverables WHERE id=?", did)
    if not row:
        raise HTTPException(404)
    db.update("deliverables", did, share_enabled=0)
    return {"ok": True}


@router.get("/hook/{token}/d/{share_id}")
async def public_share(token: str, share_id: str):
    if not hmac.compare_digest(token, _hook_token()):
        raise HTTPException(401)
    row = db.one("SELECT * FROM deliverables WHERE share_id=? AND share_enabled=1", share_id)
    if not row:
        raise HTTPException(404, "not shared (or revoked)")
    return _render_preview(row)


# ---------------- tool registration ----------------
register_tool("attach", fn(
    "attach", "Attach an existing file from your computer (home or shared/) to your reply as a "
    "downloadable deliverable card.", {"path": S, "title": S, "kind": S}, ["path"]), _attach,
    policy="allow")

register_tool("make_document", fn(
    "make_document", "Create a polished document (report, memo, brief) from Markdown and attach "
    "it. format='docx' (Word, default), 'pdf', or 'md'.",
    {"title": S, "markdown": S, "format": {"type": "string", "enum": ["docx", "pdf", "md"]}},
    ["title", "markdown"]), _make_document, policy="allow")

register_tool("make_spreadsheet", fn(
    "make_spreadsheet", "Create a spreadsheet from tabular data and attach it. sheets=[{name, "
    "columns:[...], rows:[[...],...]}]. format='xlsx' (default) or 'csv'.",
    {"title": S, "sheets": {"type": "array", "items": {"type": "object", "properties": {
        "name": S, "columns": {"type": "array", "items": S},
        "rows": {"type": "array", "items": {"type": "array"}}}}},
     "format": {"type": "string", "enum": ["xlsx", "csv"]}}, ["title", "sheets"]),
    _make_spreadsheet, policy="allow")

register_tool("make_slides", fn(
    "make_slides", "Create a clean pastel slide deck and attach it (.pptx). slides=[{title, "
    "bullets:[...], notes?, image_path?}].",
    {"title": S, "slides": {"type": "array", "items": {"type": "object", "properties": {
        "title": S, "bullets": {"type": "array", "items": S}, "notes": S,
        "image_path": S}}}}, ["title", "slides"]), _make_slides, policy="allow")

register_tool("make_chart", fn(
    "make_chart", "Draw a chart (bar/line/pie) inside your reply. "
    "data={labels:[...], series:[{name, values:[...]}], y?:{min, max}}. The axis fits the "
    "data by itself (bars start at 0); set y only to make it reach a value that matters, "
    "like a threshold.",
    {"title": S, "type": {"type": "string", "enum": ["bar", "line", "pie"]},
     "data": {"type": "object", "properties": {}}}, ["title", "type", "data"]), _make_chart,
    policy="allow")

register_tool("publish_page", CORE_TOOLS["publish_page"], _publish_page, policy="allow")
